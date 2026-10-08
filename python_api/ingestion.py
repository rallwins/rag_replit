from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath

import tiktoken
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from pypdf import PdfReader

from .config import get_settings


@dataclass(frozen=True)
class SourcePage:
    text: str
    page_number: int | None


@dataclass(frozen=True)
class Chunk:
    content: str
    token_count: int
    page_number: int | None
    start_offset: int
    end_offset: int


class DocumentExtractionError(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def extract_pages(data: bytes, filename: str, media_type: str) -> list[SourcePage]:
    suffix = PurePath(filename).suffix.lower()
    normalized_type = media_type.lower().split(";")[0].strip()
    if normalized_type == "application/pdf" or suffix == ".pdf":
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise DocumentExtractionError("encrypted_document")
            if len(reader.pages) > 1000:
                raise DocumentExtractionError("page_limit_exceeded")
            pages = [
                SourcePage(text=(page.extract_text() or "").replace("\x00", "").strip(), page_number=index)
                for index, page in enumerate(reader.pages, start=1)
            ]
        except DocumentExtractionError:
            raise
        except Exception as error:
            raise DocumentExtractionError("invalid_pdf") from error
    elif (
        normalized_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == ".docx"
    ):
        try:
            document = WordDocument(BytesIO(data))
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            if table_rows:
                content = "\n".join(part for part in (content, "\n".join(table_rows)) if part)
            pages = [SourcePage(text=content.replace("\x00", "").strip(), page_number=None)]
        except Exception as error:
            raise DocumentExtractionError("invalid_docx") from error
    elif normalized_type == "text/html" or suffix in {".html", ".htm"}:
        try:
            soup = BeautifulSoup(data.decode("utf-8-sig"), "html.parser")
            for node in soup(["script", "style", "noscript"]):
                node.decompose()
            pages = [SourcePage(text=soup.get_text("\n", strip=True), page_number=None)]
        except UnicodeDecodeError as error:
            raise DocumentExtractionError("invalid_text_encoding") from error
    elif normalized_type in {"text/plain", "text/markdown"} or suffix in {".txt", ".md", ".markdown"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise DocumentExtractionError("invalid_text_encoding") from error
        pages = [SourcePage(text=text.replace("\x00", "").strip(), page_number=None)]
    else:
        raise DocumentExtractionError("unsupported_media_type")

    if not pages or not any(page.text.strip() for page in pages):
        raise DocumentExtractionError("no_extractable_text")
    return pages


def chunk_pages(pages: list[SourcePage]) -> list[Chunk]:
    settings = get_settings()
    encoder = tiktoken.get_encoding("cl100k_base")
    step = settings.chunk_size_tokens - settings.chunk_overlap_tokens
    chunks: list[Chunk] = []
    for page in pages:
        tokens = encoder.encode(page.text)
        if not tokens:
            continue
        start = 0
        while start < len(tokens):
            end = min(start + settings.chunk_size_tokens, len(tokens))
            content = encoder.decode(tokens[start:end]).strip()
            if content:
                char_start = len(encoder.decode(tokens[:start]))
                char_end = len(encoder.decode(tokens[:end]))
                chunks.append(
                    Chunk(
                        content=content,
                        token_count=end - start,
                        page_number=page.page_number,
                        start_offset=char_start,
                        end_offset=char_end,
                    )
                )
            if end == len(tokens):
                break
            start += step
            if len(chunks) > 20_000:
                raise DocumentExtractionError("chunk_limit_exceeded")
    return chunks
